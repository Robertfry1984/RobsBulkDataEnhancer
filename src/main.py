import csv
import json
import os
import socket
import ipaddress
import sys
import time
from dataclasses import dataclass
from typing import List, Optional
from urllib.parse import urlparse

import keyring
import requests
from PyQt6.QtCore import (
    QAbstractTableModel,
    QModelIndex,
    QObject,
    Qt,
    QSettings,
    QThread,
    pyqtSignal,
)
from PyQt6.QtGui import QAction, QFont, QKeySequence
from PyQt6.QtWidgets import (
    QApplication,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QPlainTextEdit,
    QDoubleSpinBox,
    QSpinBox,
    QSplitter,
    QStatusBar,
    QTableView,
    QCheckBox,
    QVBoxLayout,
    QWidget,
)

APP_ORG = "Robs"
APP_NAME = "RobsBulkDataEnhancer"
KEYRING_SERVICE = "RobsBulkDataEnhancer"

ROOT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
VENDOR_DIR = os.path.join(ROOT_DIR, "vendor")
if os.path.isdir(VENDOR_DIR):
    for name in ["ddgs", "trafilatura"]:
        path = os.path.join(VENDOR_DIR, name)
        if os.path.isdir(path):
            sys.path.insert(0, path)


def column_label(index: int) -> str:
    label = ""
    idx = index
    while True:
        idx, rem = divmod(idx, 26)
        label = chr(65 + rem) + label
        if idx == 0:
            break
        idx -= 1
    return label


def column_index(label: str) -> Optional[int]:
    if not label:
        return None
    text = label.strip().upper()
    if text.isdigit():
        value = int(text) - 1
        return value if value >= 0 else None
    total = 0
    for ch in text:
        if not ("A" <= ch <= "Z"):
            return None
        total = total * 26 + (ord(ch) - 64)
    return total - 1


def extract_json_path(data, path: str):
    current = data
    for part in path.split("."):
        if part.isdigit():
            idx = int(part)
            current = current[idx]
        else:
            current = current[part]
    return current


def sanitize_clipboard_cell(text: str) -> str:
    if text is None:
        return ""
    value = str(text)
    value = value.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    value = value.replace("\t", " ")
    return value


def is_public_host(hostname: str) -> bool:
    if not hostname:
        return False
    try:
        if hostname.lower() in {"localhost"}:
            return False
        ip = ipaddress.ip_address(hostname)
        return not (
            ip.is_private
            or ip.is_loopback
            or ip.is_link_local
            or ip.is_multicast
            or ip.is_reserved
            or ip.is_unspecified
        )
    except ValueError:
        try:
            infos = socket.getaddrinfo(hostname, None)
        except socket.gaierror:
            return False
        for info in infos:
            addr = info[4][0]
            try:
                ip = ipaddress.ip_address(addr)
            except ValueError:
                return False
            if (
                ip.is_private
                or ip.is_loopback
                or ip.is_link_local
                or ip.is_multicast
                or ip.is_reserved
                or ip.is_unspecified
            ):
                return False
        return True


def is_allowed_web_url(url: str) -> bool:
    try:
        parsed = urlparse(url)
    except Exception:
        return False
    if parsed.scheme not in {"http", "https"}:
        return False
    return is_public_host(parsed.hostname or "")


class DataTableModel(QAbstractTableModel):
    def __init__(self, rows: int = 100, cols: int = 10, parent: Optional[QObject] = None):
        super().__init__(parent)
        self._data: List[List[str]] = [["" for _ in range(cols)] for _ in range(rows)]

    def rowCount(self, parent=QModelIndex()) -> int:
        return len(self._data)

    def columnCount(self, parent=QModelIndex()) -> int:
        return len(self._data[0]) if self._data else 0

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.EditRole):
            return self._data[index.row()][index.column()]
        return None

    def setData(self, index, value, role=Qt.ItemDataRole.EditRole) -> bool:
        if role != Qt.ItemDataRole.EditRole or not index.isValid():
            return False
        self._data[index.row()][index.column()] = str(value)
        self.dataChanged.emit(index, index, [Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.EditRole])
        return True

    def flags(self, index):
        if not index.isValid():
            return Qt.ItemFlag.ItemIsEnabled
        return (
            Qt.ItemFlag.ItemIsSelectable
            | Qt.ItemFlag.ItemIsEnabled
            | Qt.ItemFlag.ItemIsEditable
        )

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role != Qt.ItemDataRole.DisplayRole:
            return None
        if orientation == Qt.Orientation.Horizontal:
            return column_label(section)
        return str(section + 1)

    def ensure_size(self, rows: int, cols: int):
        if rows <= self.rowCount() and cols <= self.columnCount():
            return
        new_rows = max(rows, self.rowCount())
        new_cols = max(cols, self.columnCount())
        self.beginResetModel()
        for row in self._data:
            if len(row) < new_cols:
                row.extend([""] * (new_cols - len(row)))
        if len(self._data) < new_rows:
            for _ in range(new_rows - len(self._data)):
                self._data.append([""] * new_cols)
        self.endResetModel()

    def clear(self):
        self.beginResetModel()
        self._data = [[""] * self.columnCount() for _ in range(self.rowCount())]
        self.endResetModel()

    def set_block(self, start_row: int, start_col: int, values: List[List[str]]):
        if not values:
            return
        rows = start_row + len(values)
        cols = start_col + max(len(r) for r in values)
        self.ensure_size(rows, cols)
        for r, row_vals in enumerate(values):
            for c, val in enumerate(row_vals):
                self._data[start_row + r][start_col + c] = str(val)
        top_left = self.index(start_row, start_col)
        bottom_right = self.index(rows - 1, cols - 1)
        self.dataChanged.emit(top_left, bottom_right, [Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.EditRole])

    def set_cell(self, row: int, col: int, value: str):
        self.ensure_size(row + 1, col + 1)
        self._data[row][col] = value
        idx = self.index(row, col)
        self.dataChanged.emit(idx, idx, [Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.EditRole])

    def get_cell(self, row: int, col: int) -> str:
        if row < 0 or col < 0:
            return ""
        if row >= self.rowCount() or col >= self.columnCount():
            return ""
        return self._data[row][col]

    def to_rows(self) -> List[List[str]]:
        return self._data


class DataTableView(QTableView):
    copyRequested = pyqtSignal()
    pasteRequested = pyqtSignal()

    def keyPressEvent(self, event):
        if event.matches(QKeySequence.StandardKey.Copy):
            self.copyRequested.emit()
            return
        if event.matches(QKeySequence.StandardKey.Paste):
            self.pasteRequested.emit()
            return
        super().keyPressEvent(event)


@dataclass
class AppSettings:
    endpoint: str
    auth_header: str
    request_template: str
    response_json_path: str
    timeout_seconds: int
    model: str
    temperature: float
    theme: str


class SettingsDialog(QDialog):
    def __init__(self, settings: AppSettings, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Settings")
        self.setMinimumWidth(520)

        self.endpoint_input = QLineEdit(settings.endpoint)
        self.auth_header_input = QLineEdit(settings.auth_header)
        self.auth_key_input = QLineEdit()
        self.auth_key_input.setEchoMode(QLineEdit.EchoMode.Password)
        self.request_template_input = QPlainTextEdit(settings.request_template)
        self.response_path_input = QLineEdit(settings.response_json_path)
        self.timeout_input = QSpinBox()
        self.timeout_input.setRange(1, 120)
        self.timeout_input.setValue(settings.timeout_seconds)
        self.model_input = QLineEdit(settings.model)
        self.temperature_input = QDoubleSpinBox()
        self.temperature_input.setRange(0.0, 2.0)
        self.temperature_input.setDecimals(2)
        self.temperature_input.setSingleStep(0.1)
        self.temperature_input.setValue(settings.temperature)

        form = QFormLayout()
        form.addRow("Endpoint URL", self.endpoint_input)
        form.addRow("Auth header name", self.auth_header_input)
        form.addRow("Auth key", self.auth_key_input)
        form.addRow("Request template (JSON)", self.request_template_input)
        form.addRow("Response JSON path", self.response_path_input)
        form.addRow("Timeout (seconds)", self.timeout_input)
        form.addRow("Model", self.model_input)
        form.addRow("Temperature", self.temperature_input)

        help_label = QLabel(
            "Template supports {prompt_json}, {prompt}, {system_json}, {system}, {model}, {model_json}, {temperature}."
        )
        help_label.setWordWrap(True)

        buttons = QHBoxLayout()
        save_btn = QPushButton("Save")
        cancel_btn = QPushButton("Cancel")
        save_btn.clicked.connect(self.accept)
        cancel_btn.clicked.connect(self.reject)
        buttons.addStretch(1)
        buttons.addWidget(save_btn)
        buttons.addWidget(cancel_btn)

        layout = QVBoxLayout()
        layout.addLayout(form)
        layout.addWidget(help_label)
        layout.addLayout(buttons)
        self.setLayout(layout)

    def get_values(self):
        return {
            "endpoint": self.endpoint_input.text().strip(),
            "auth_header": self.auth_header_input.text().strip(),
            "auth_key": self.auth_key_input.text(),
            "request_template": self.request_template_input.toPlainText().strip(),
            "response_json_path": self.response_path_input.text().strip(),
            "timeout_seconds": int(self.timeout_input.value()),
            "model": self.model_input.text().strip(),
            "temperature": float(self.temperature_input.value()),
        }


class ProcessingWorker(QThread):
    cellReady = pyqtSignal(int, int, str)
    progress = pyqtSignal(int, int)
    error = pyqtSignal(str)
    finished = pyqtSignal()
    stopped = pyqtSignal()

    def __init__(
        self,
        mode: str,
        rows: List[int],
        input_values: Optional[List[str]],
        output_col: int,
        question: str,
        response_hint: str,
        delay_ms: int,
        settings: AppSettings,
        websearch_enabled: bool,
        max_sources: int,
        max_pages: int,
        max_extract_chars: int,
        enforce_format: bool,
    ):
        super().__init__()
        self._stop_requested = False
        self.mode = mode
        self.rows = rows
        self.input_values = input_values or []
        self.output_col = output_col
        self.question = question
        self.response_hint = response_hint
        self.delay_ms = delay_ms
        self.settings = settings
        self.websearch_enabled = websearch_enabled
        self.max_sources = max_sources
        self.max_pages = max_pages
        self.max_extract_chars = max_extract_chars
        self.enforce_format = enforce_format

    def stop(self):
        self._stop_requested = True

    def build_prompt(self, input_value: str, index_hint: str, include_response_hint: bool) -> str:
        parts = []
        if self.question:
            parts.append(self.question)
        if input_value:
            parts.append(f"Input: {input_value}")
        if include_response_hint and self.response_hint:
            parts.append(f"Response format: {self.response_hint}")
        if index_hint:
            parts.append(index_hint)
        return "\n".join(parts).strip()

    def build_payload(self, prompt: str, system_text: str):
        template = self.settings.request_template.strip() if self.settings.request_template else ""
        if (not template or "\"messages\"" not in template) and "/v1/chat/completions" in self.settings.endpoint:
            template = (
                '{"model": "{model}", "messages": '
                '[{"role": "system", "content": {system_json}}, '
                '{"role": "user", "content": {prompt_json}}], '
                '"temperature": {temperature}}'
            )
        if template:
            rendered = template
            rendered = rendered.replace("{prompt_json}", json.dumps(prompt))
            rendered = rendered.replace("{prompt}", prompt)
            rendered = rendered.replace("{system_json}", json.dumps(system_text))
            rendered = rendered.replace("{system}", system_text)
            rendered = rendered.replace("{model_json}", json.dumps(self.settings.model))
            rendered = rendered.replace("{model}", self.settings.model)
            rendered = rendered.replace("{temperature}", str(self.settings.temperature))
            try:
                return json.loads(rendered)
            except json.JSONDecodeError:
                return {"prompt": prompt}
        return {"prompt": prompt}

    def extract_response_text(self, response: requests.Response) -> str:
        if self.settings.response_json_path:
            try:
                data = response.json()
                value = extract_json_path(data, self.settings.response_json_path)
                return str(value).strip()
            except Exception:
                return response.text.strip()
        return response.text.strip()

    def infer_word_limit(self, hint: str) -> Optional[int]:
        text = hint.lower()
        if "1 or 2 word" in text or "one or two word" in text or "1-2 word" in text:
            return 2
        if "two word" in text or "2 word" in text:
            return 2
        if "one word" in text or "1 word" in text or "single word" in text:
            return 1
        if "single" in text and "word" in text:
            return 1
        return None

    def enforce_response_format(self, text: str, hint: str) -> str:
        cleaned = text.strip()
        if not cleaned:
            return cleaned
        cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
        cleaned = cleaned.split("\n", 1)[0]
        cleaned = cleaned.lstrip("-*•0123456789. )(").strip()

        lower_hint = hint.lower()
        if "sentence" in lower_hint:
            for end in [". ", "! ", "? "]:
                if end in cleaned:
                    cleaned = cleaned.split(end, 1)[0].strip() + end.strip()
                    break

        word_limit = self.infer_word_limit(hint)
        if word_limit:
            words = cleaned.split()
            if len(words) > word_limit:
                cleaned = " ".join(words[:word_limit])

        return cleaned.strip()

    def search_sources(self, query: str) -> List[dict]:
        try:
            from ddgs import DDGS
        except Exception as exc:
            self.error.emit(f"Web search unavailable (ddgs import failed): {exc}")
            return []

        results = []
        try:
            with DDGS() as ddgs:
                for item in ddgs.text(
                    query, max_results=self.max_pages, safesearch="moderate"
                ):
                    url = item.get("href") or item.get("url")
                    if url and is_allowed_web_url(url):
                        results.append(
                            {
                                "title": item.get("title", ""),
                                "url": url,
                            }
                        )
        except Exception as exc:
            self.error.emit(f"Web search failed: {exc}")
        return results

    def fetch_and_extract(self, url: str) -> str:
        try:
            import trafilatura
        except Exception as exc:
            self.error.emit(f"Web extraction unavailable (trafilatura import failed): {exc}")
            return ""

        try:
            if not is_allowed_web_url(url):
                return ""
            response = requests.get(
                url,
                timeout=(5, 10),
                headers={"User-Agent": "RobsBulkDataEnhancer/1.0"},
                stream=True,
            )
            response.raise_for_status()
            max_bytes = 1_000_000
            chunks = []
            total = 0
            for chunk in response.iter_content(chunk_size=8192):
                if not chunk:
                    continue
                total += len(chunk)
                if total > max_bytes:
                    break
                chunks.append(chunk)
            html = b"".join(chunks).decode(response.encoding or "utf-8", errors="ignore")
            text = trafilatura.extract(
                html, include_comments=False, include_tables=False, favor_recall=True
            )
            return text or ""
        except Exception:
            return ""

    def build_web_context(self, query: str) -> str:
        sources = self.search_sources(query)
        if not sources:
            return ""
        chunks = []
        for source in sources[: self.max_sources]:
            content = self.fetch_and_extract(source["url"])
            if not content:
                continue
            content = content.strip()
            if len(content) > self.max_extract_chars:
                content = content[: self.max_extract_chars] + "..."
            title = source.get("title", "")
            url = source.get("url", "")
            chunks.append(f"Source: {title}\nURL: {url}\nContent: {content}")
        return "\n\n".join(chunks)

    def run(self):
        if not self.settings.endpoint:
            self.error.emit("No endpoint configured. Open Settings to set the endpoint.")
            self.finished.emit()
            return
        parsed_endpoint = urlparse(self.settings.endpoint)
        if parsed_endpoint.scheme not in {"http", "https"}:
            self.error.emit("Endpoint must be http or https.")
            self.finished.emit()
            return

        headers = {}
        if self.settings.auth_header:
            auth_key = keyring.get_password(KEYRING_SERVICE, "api_key")
            if auth_key:
                headers[self.settings.auth_header] = auth_key

        session = requests.Session()
        total = len(self.rows)

        for idx, row in enumerate(self.rows, start=1):
            if self._stop_requested:
                self.stopped.emit()
                return

            input_value = ""
            if self.mode == "range" and idx - 1 < len(self.input_values):
                input_value = self.input_values[idx - 1]

            index_hint = f"Item {idx} of {total}" if self.mode == "generate" else ""
            system_text = self.response_hint.strip()
            include_hint = not ("/v1/chat/completions" in self.settings.endpoint)

            if self.websearch_enabled:
                base_query = " ".join([p for p in [self.question, input_value] if p])
                web_context = self.build_web_context(base_query)
                prompt_parts = []
                if self.question:
                    prompt_parts.append(self.question)
                if input_value:
                    prompt_parts.append(f"Input: {input_value}")
                if web_context:
                    prompt_parts.append(
                        "Use the sources below to answer. Keep the response concise.\n"
                        + web_context
                    )
                if index_hint:
                    prompt_parts.append(index_hint)
                prompt = "\n".join(prompt_parts).strip()
            else:
                prompt = self.build_prompt(input_value, index_hint, include_hint)

            payload = self.build_payload(prompt, system_text)

            try:
                response = session.post(
                    self.settings.endpoint,
                    json=payload,
                    headers=headers,
                    timeout=self.settings.timeout_seconds,
                )
                response.raise_for_status()
                text = self.extract_response_text(response)
                if self.enforce_format and self.response_hint:
                    text = self.enforce_response_format(text, self.response_hint)
            except Exception as exc:
                self.error.emit(str(exc))
                text = ""

            self.cellReady.emit(row, self.output_col, text)
            self.progress.emit(idx, total)

            if self.delay_ms > 0 and idx < total:
                time.sleep(self.delay_ms / 1000.0)

        self.finished.emit()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Robs Bulk Data Enhancer")
        self.resize(1400, 800)

        self.settings_store = QSettings(APP_ORG, APP_NAME)
        self.app_settings = self.load_settings()

        self.model = DataTableModel()
        self.worker: Optional[ProcessingWorker] = None

        self.init_ui()
        self.apply_theme(self.app_settings.theme)

    def init_ui(self):
        self.status = QStatusBar()
        self.setStatusBar(self.status)

        self.table = DataTableView()
        self.table.setModel(self.model)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.verticalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Fixed)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionMode(QTableView.SelectionMode.ExtendedSelection)
        self.table.setSelectionBehavior(QTableView.SelectionBehavior.SelectItems)
        self.table.copyRequested.connect(self.copy_selection)
        self.table.pasteRequested.connect(self.paste_clipboard)

        self.question_input = QPlainTextEdit()
        self.question_input.setPlaceholderText("Type the main question or prompt here...")

        self.input_col = QLineEdit("A")
        self.output_col = QLineEdit("B")
        self.row_start = QSpinBox()
        self.row_start.setRange(1, 2_000_000)
        self.row_start.setValue(1)
        self.row_end = QSpinBox()
        self.row_end.setRange(1, 2_000_000)
        self.row_end.setValue(100)
        self.response_hint = QPlainTextEdit()
        self.response_hint.setPlaceholderText("Describe the response format, e.g. 'Return one category only'.")
        self.delay_ms = QSpinBox()
        self.delay_ms.setRange(0, 10_000)
        self.delay_ms.setValue(100)
        self.websearch_toggle = QCheckBox("Use web search")
        self.websearch_toggle.setChecked(False)
        self.enforce_format_toggle = QCheckBox("Enforce response format")
        self.enforce_format_toggle.setChecked(True)
        self.max_sources = QSpinBox()
        self.max_sources.setRange(1, 10)
        self.max_sources.setValue(3)
        self.max_pages = QSpinBox()
        self.max_pages.setRange(1, 20)
        self.max_pages.setValue(5)
        self.max_extract_chars = QSpinBox()
        self.max_extract_chars.setRange(200, 10000)
        self.max_extract_chars.setSingleStep(200)
        self.max_extract_chars.setValue(2000)

        self.generate_count = QSpinBox()
        self.generate_count.setRange(1, 1_000_000)
        self.generate_count.setValue(100)

        self.seed_prefix = QLineEdit("Item ")
        self.seed_count = QSpinBox()
        self.seed_count.setRange(1, 1_000_000)
        self.seed_count.setValue(100)

        go_btn = QPushButton("Go")
        stop_btn = QPushButton("Stop")
        gen_btn = QPushButton("Generate Rows")
        fill_btn = QPushButton("Fill Column A")
        go_btn.clicked.connect(self.start_processing)
        stop_btn.clicked.connect(self.stop_processing)
        gen_btn.clicked.connect(self.generate_rows)
        fill_btn.clicked.connect(self.fill_column_a)

        right_form = QFormLayout()
        right_form.addRow("Input column", self.input_col)
        right_form.addRow("Output column", self.output_col)
        right_form.addRow("Row start", self.row_start)
        right_form.addRow("Row end", self.row_end)
        right_form.addRow("Response format", self.response_hint)
        right_form.addRow("Delay (ms)", self.delay_ms)
        right_form.addRow("Web search", self.websearch_toggle)
        right_form.addRow("Enforce format", self.enforce_format_toggle)
        right_form.addRow("Max sources", self.max_sources)
        right_form.addRow("Max pages per query", self.max_pages)
        right_form.addRow("Max extract chars", self.max_extract_chars)
        right_form.addRow(QLabel(""))
        right_form.addRow(QLabel("Generate rows (no input column)"))
        right_form.addRow("Count", self.generate_count)
        right_form.addRow(QLabel("Seed list in Column A"))
        right_form.addRow("Prefix", self.seed_prefix)
        right_form.addRow("Count", self.seed_count)

        right_buttons = QHBoxLayout()
        right_buttons.addWidget(go_btn)
        right_buttons.addWidget(stop_btn)
        right_buttons.addWidget(gen_btn)
        right_buttons.addWidget(fill_btn)

        right_panel = QWidget()
        right_layout = QVBoxLayout()
        right_layout.addLayout(right_form)
        right_layout.addLayout(right_buttons)
        right_layout.addStretch(1)
        right_panel.setLayout(right_layout)

        left_panel = QWidget()
        left_layout = QVBoxLayout()
        left_layout.addWidget(QLabel("Main question"))
        left_layout.addWidget(self.question_input)
        left_panel.setLayout(left_layout)

        splitter = QSplitter()
        splitter.addWidget(left_panel)
        splitter.addWidget(self.table)
        splitter.addWidget(right_panel)
        splitter.setStretchFactor(0, 1)
        splitter.setStretchFactor(1, 4)
        splitter.setStretchFactor(2, 2)
        splitter.setSizes([260, 780, 360])

        container = QWidget()
        container_layout = QVBoxLayout()
        container_layout.addWidget(splitter)
        container.setLayout(container_layout)
        self.setCentralWidget(container)

        self.init_menu()

    def init_menu(self):
        menu = self.menuBar()

        file_menu = menu.addMenu("File")
        open_action = QAction("Open...", self)
        save_action = QAction("Save As...", self)
        clear_action = QAction("Clear", self)
        open_action.triggered.connect(self.open_file)
        save_action.triggered.connect(self.save_as)
        clear_action.triggered.connect(self.clear_data)
        file_menu.addAction(open_action)
        file_menu.addAction(save_action)
        file_menu.addSeparator()
        file_menu.addAction(clear_action)

        settings_menu = menu.addMenu("Settings")
        open_settings = QAction("API Settings...", self)
        open_settings.triggered.connect(self.open_settings_dialog)
        settings_menu.addAction(open_settings)

        view_menu = menu.addMenu("View")
        light_action = QAction("Light Mode", self)
        dark_action = QAction("Dark Mode", self)
        light_action.triggered.connect(lambda: self.apply_theme("light"))
        dark_action.triggered.connect(lambda: self.apply_theme("dark"))
        view_menu.addAction(light_action)
        view_menu.addAction(dark_action)

        help_menu = menu.addMenu("Help")
        about_action = QAction("About", self)
        about_action.triggered.connect(self.show_about)
        help_menu.addAction(about_action)

    def show_about(self):
        QMessageBox.information(
            self,
            "About",
            "Robs Bulk Data Enhancer\nA data list creation & enhancement tool.",
        )

    def load_settings(self) -> AppSettings:
        endpoint = self.settings_store.value("endpoint", "")
        auth_header = self.settings_store.value("auth_header", "Authorization")
        request_template = self.settings_store.value(
            "request_template",
            '{"model":"{model}","messages":[{"role":"system","content":{system_json}},{"role":"user","content":{prompt_json}}],"temperature":{temperature}}',
        )
        response_json_path = self.settings_store.value("response_json_path", "")
        if not response_json_path:
            response_json_path = "choices.0.message.content"
        timeout_seconds = int(self.settings_store.value("timeout_seconds", 30))
        model = self.settings_store.value("model", "gpt-4o-mini")
        temperature = float(self.settings_store.value("temperature", 0.2))
        theme = self.settings_store.value("theme", "light")
        return AppSettings(
            endpoint=endpoint,
            auth_header=auth_header,
            request_template=request_template,
            response_json_path=response_json_path,
            timeout_seconds=timeout_seconds,
            model=model,
            temperature=temperature,
            theme=theme,
        )

    def save_settings(self):
        self.settings_store.setValue("endpoint", self.app_settings.endpoint)
        self.settings_store.setValue("auth_header", self.app_settings.auth_header)
        self.settings_store.setValue("request_template", self.app_settings.request_template)
        self.settings_store.setValue("response_json_path", self.app_settings.response_json_path)
        self.settings_store.setValue("timeout_seconds", self.app_settings.timeout_seconds)
        self.settings_store.setValue("model", self.app_settings.model)
        self.settings_store.setValue("temperature", self.app_settings.temperature)
        self.settings_store.setValue("theme", self.app_settings.theme)

    def open_settings_dialog(self):
        dialog = SettingsDialog(self.app_settings, self)
        if dialog.exec() == QDialog.DialogCode.Accepted:
            values = dialog.get_values()
            self.app_settings.endpoint = values["endpoint"]
            self.app_settings.auth_header = values["auth_header"]
            self.app_settings.request_template = values["request_template"]
            self.app_settings.response_json_path = values["response_json_path"]
            self.app_settings.timeout_seconds = values["timeout_seconds"]
            self.app_settings.model = values["model"]
            self.app_settings.temperature = values["temperature"]
            if values["auth_key"]:
                keyring.set_password(KEYRING_SERVICE, "api_key", values["auth_key"])
            self.save_settings()
            self.status.showMessage("Settings saved.", 3000)

    def apply_theme(self, theme: str):
        base_font = QFont("Segoe UI", 10)
        self.setFont(base_font)

        if theme == "dark":
            self.setStyleSheet(
                """
                QMainWindow { background: #1b1f24; color: #e5e9f0; }
                QWidget { color: #e5e9f0; }
                QPlainTextEdit, QLineEdit, QTableView { background: #222831; color: #e5e9f0; border: 1px solid #2f3943; }
                QHeaderView::section { background: #2c343d; color: #e5e9f0; padding: 4px; border: 1px solid #2f3943; }
                QPushButton { background: #3a6ea5; color: #ffffff; border-radius: 6px; padding: 6px 10px; }
                QPushButton:hover { background: #4b7fb6; }
                QMenuBar { background: #1b1f24; }
                QMenu { background: #222831; }
                QStatusBar { background: #1b1f24; }
                """
            )
        else:
            self.setStyleSheet(
                """
                QMainWindow { background: #f4f6f8; color: #1e2125; }
                QWidget { color: #1e2125; }
                QPlainTextEdit, QLineEdit, QTableView { background: #ffffff; color: #1e2125; border: 1px solid #d7dce2; }
                QHeaderView::section { background: #eef1f4; color: #1e2125; padding: 4px; border: 1px solid #d7dce2; }
                QPushButton { background: #2f7ed8; color: #ffffff; border-radius: 6px; padding: 6px 10px; }
                QPushButton:hover { background: #3f8ee8; }
                QMenuBar { background: #f4f6f8; }
                QMenu { background: #ffffff; }
                QStatusBar { background: #f4f6f8; }
                """
            )
        self.app_settings.theme = theme
        self.save_settings()

    def copy_selection(self):
        selection = self.table.selectionModel().selectedIndexes()
        if not selection:
            return
        rows = [idx.row() for idx in selection]
        cols = [idx.column() for idx in selection]
        min_row, max_row = min(rows), max(rows)
        min_col, max_col = min(cols), max(cols)

        data_lines = []
        for r in range(min_row, max_row + 1):
            row_values = []
            for c in range(min_col, max_col + 1):
                row_values.append(sanitize_clipboard_cell(self.model.get_cell(r, c)))
            data_lines.append("\t".join(row_values))
        QApplication.clipboard().setText("\r\n".join(data_lines))

    def paste_clipboard(self):
        text = QApplication.clipboard().text()
        if not text:
            return
        rows = [row.split("\t") for row in text.splitlines()]
        selection = self.table.selectionModel().selectedIndexes()
        if selection:
            start_row = min(idx.row() for idx in selection)
            start_col = min(idx.column() for idx in selection)
        else:
            start_row, start_col = 0, 0
        self.model.set_block(start_row, start_col, rows)

    def parse_column(self, text: str) -> Optional[int]:
        return column_index(text)

    def validate_range(self, start: int, end: int) -> Optional[List[int]]:
        if start <= 0 or end <= 0 or end < start:
            return None
        return list(range(start - 1, end))

    def start_processing(self):
        if self.worker and self.worker.isRunning():
            QMessageBox.warning(self, "Processing", "Processing is already running.")
            return
        input_col = self.parse_column(self.input_col.text())
        output_col = self.parse_column(self.output_col.text())
        if input_col is None or output_col is None:
            QMessageBox.warning(self, "Input", "Invalid column selection.")
            return
        rows = self.validate_range(self.row_start.value(), self.row_end.value())
        if rows is None:
            QMessageBox.warning(self, "Input", "Invalid row range.")
            return
        input_values = [self.model.get_cell(r, input_col) for r in rows]

        self.worker = ProcessingWorker(
            mode="range",
            rows=rows,
            input_values=input_values,
            output_col=output_col,
            question=self.question_input.toPlainText().strip(),
            response_hint=self.response_hint.toPlainText().strip(),
            delay_ms=int(self.delay_ms.value()),
            settings=self.app_settings,
            websearch_enabled=self.websearch_toggle.isChecked(),
            max_sources=int(self.max_sources.value()),
            max_pages=int(self.max_pages.value()),
            max_extract_chars=int(self.max_extract_chars.value()),
            enforce_format=self.enforce_format_toggle.isChecked(),
        )
        self.worker.setParent(self)
        self.worker.cellReady.connect(self.model.set_cell)
        self.worker.progress.connect(self.update_progress)
        self.worker.error.connect(self.show_error)
        self.worker.finished.connect(lambda: self.status.showMessage("Finished", 3000))
        self.worker.stopped.connect(lambda: self.status.showMessage("Stopped", 3000))
        self.worker.start()
        self.status.showMessage("Processing...")

    def generate_rows(self):
        if self.worker and self.worker.isRunning():
            QMessageBox.warning(self, "Processing", "Processing is already running.")
            return
        output_col = self.parse_column(self.output_col.text())
        if output_col is None:
            QMessageBox.warning(self, "Input", "Invalid output column.")
            return
        count = int(self.generate_count.value())
        rows = list(range(0, count))

        self.worker = ProcessingWorker(
            mode="generate",
            rows=rows,
            input_values=None,
            output_col=output_col,
            question=self.question_input.toPlainText().strip(),
            response_hint=self.response_hint.toPlainText().strip(),
            delay_ms=int(self.delay_ms.value()),
            settings=self.app_settings,
            websearch_enabled=self.websearch_toggle.isChecked(),
            max_sources=int(self.max_sources.value()),
            max_pages=int(self.max_pages.value()),
            max_extract_chars=int(self.max_extract_chars.value()),
            enforce_format=self.enforce_format_toggle.isChecked(),
        )
        self.worker.setParent(self)
        self.worker.cellReady.connect(self.model.set_cell)
        self.worker.progress.connect(self.update_progress)
        self.worker.error.connect(self.show_error)
        self.worker.finished.connect(lambda: self.status.showMessage("Finished", 3000))
        self.worker.stopped.connect(lambda: self.status.showMessage("Stopped", 3000))
        self.worker.start()
        self.status.showMessage("Generating...")

    def stop_processing(self):
        if self.worker and self.worker.isRunning():
            self.worker.stop()
            self.status.showMessage("Stopping...", 2000)

    def update_progress(self, current: int, total: int):
        self.status.showMessage(f"Processing {current}/{total}")

    def show_error(self, message: str):
        self.status.showMessage("Error occurred", 4000)
        QMessageBox.warning(self, "Error", message)

    def fill_column_a(self):
        count = int(self.seed_count.value())
        prefix = self.seed_prefix.text()
        values = [[f"{prefix}{i + 1}"] for i in range(count)]
        self.model.set_block(0, 0, values)

    def clear_data(self):
        self.model.clear()

    def save_as(self):
        path, filter_name = QFileDialog.getSaveFileName(
            self,
            "Save As",
            "",
            "CSV Files (*.csv);;Excel Files (*.xlsx)",
        )
        if not path:
            return
        if path.lower().endswith(".xlsx"):
            self.save_xlsx(path)
        else:
            if not path.lower().endswith(".csv"):
                path = f"{path}.csv"
            self.save_csv(path)

    def save_csv(self, path: str):
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.writer(handle)
                writer.writerows(self.model.to_rows())
            self.status.showMessage("Saved CSV", 3000)
        except Exception as exc:
            QMessageBox.warning(self, "Save Error", str(exc))

    def save_xlsx(self, path: str):
        try:
            from openpyxl import Workbook

            wb = Workbook(write_only=True)
            ws = wb.create_sheet()
            for row in self.model.to_rows():
                ws.append(row)
            wb.save(path)
            self.status.showMessage("Saved XLSX", 3000)
        except Exception as exc:
            QMessageBox.warning(self, "Save Error", str(exc))

    def open_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Open File",
            "",
            "CSV Files (*.csv);;Excel Files (*.xlsx)",
        )
        if not path:
            return
        if path.lower().endswith(".xlsx"):
            self.load_xlsx(path)
        else:
            self.load_csv(path)

    def load_csv(self, path: str):
        try:
            with open(path, "r", newline="", encoding="utf-8-sig") as handle:
                reader = csv.reader(handle)
                rows = [row for row in reader]
            if rows:
                self.model.set_block(0, 0, rows)
            self.status.showMessage("Loaded CSV", 3000)
        except Exception as exc:
            QMessageBox.warning(self, "Open Error", str(exc))

    def load_xlsx(self, path: str):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(["" if cell is None else str(cell) for cell in row])
            if rows:
                self.model.set_block(0, 0, rows)
            self.status.showMessage("Loaded XLSX", 3000)
        except Exception as exc:
            QMessageBox.warning(self, "Open Error", str(exc))


def main():
    app = QApplication(sys.argv)
    app.setOrganizationName(APP_ORG)
    app.setApplicationName(APP_NAME)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
